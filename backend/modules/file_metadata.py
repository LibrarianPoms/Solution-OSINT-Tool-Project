# ─────────────────────────────────────────────────────────────────────────────
# file_metadata.py — File Metadata OSINT Scanner
# ─────────────────────────────────────────────────────────────────────────────
#
# Downloads a public file (PDF, DOCX, XLSX, JPG/PNG, or anything else) and
# extracts whatever embedded metadata it can find — document author/creator,
# creation & modification timestamps, software used, revision counts, and
# for images: EXIF data including camera model and (if present) embedded
# GPS coordinates. This is the classic "FOCA / Metagoofil" style OSINT
# technique — public documents often leak internal usernames, software
# versions, or device/location info their owners never meant to share.
#
# DEPENDENCIES (add to requirements.txt):
#   pip install pypdf python-docx openpyxl Pillow
#
# ─────────────────────────────────────────────────────────────────────────────

import io
import os
import requests

MAX_FILE_SIZE = 25 * 1024 * 1024  # 25MB safety cap so a huge file can't hang the server
DOWNLOAD_TIMEOUT = 15
IMAGE_EXTS = ("jpg", "jpeg", "png", "gif", "bmp", "tiff", "tif", "webp")


def _human_size(num_bytes) -> str:
    try:
        num_bytes = float(num_bytes)
    except (TypeError, ValueError):
        return "unknown"
    for unit in ["B", "KB", "MB", "GB"]:
        if num_bytes < 1024:
            return f"{int(num_bytes)} {unit}" if unit == "B" else f"{num_bytes:.1f} {unit}"
        num_bytes /= 1024
    return f"{num_bytes:.1f} TB"


def _detect_type(url: str, content_type: str, content: bytes) -> str:
    """Best-effort file type detection: magic bytes > extension > Content-Type."""
    ext = os.path.splitext(url.split("?")[0])[1].lower().lstrip(".")
    ct = (content_type or "").lower()

    if content[:4] == b"%PDF":
        return "pdf"
    if content[:4] == b"PK\x03\x04":
        if ext in ("docx", "xlsx", "pptx"):
            return ext
        if "wordprocessingml" in ct:
            return "docx"
        if "spreadsheetml" in ct:
            return "xlsx"
        if "presentationml" in ct:
            return "pptx"
        return "zip"
    if content[:3] == b"\xff\xd8\xff":
        return "jpg"
    if content[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if ext:
        return ext
    if "pdf" in ct:
        return "pdf"
    return "unknown"


def _clean(value):
    """Drop empty/placeholder values so the UI doesn't render blank rows."""
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _extract_pdf(content: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    info = reader.metadata or {}

    result = {
        "title": _clean(getattr(info, "title", None)),
        "author": _clean(getattr(info, "author", None)),
        "subject": _clean(getattr(info, "subject", None)),
        "creator": _clean(getattr(info, "creator", None)),
        "producer": _clean(getattr(info, "producer", None)),
        "created": _clean(getattr(info, "creation_date", None)),
        "modified": _clean(getattr(info, "modification_date", None)),
        "page_count": len(reader.pages),
        "encrypted": reader.is_encrypted,
    }
    return result


def _extract_docx(content: bytes) -> dict:
    from docx import Document

    doc = Document(io.BytesIO(content))
    cp = doc.core_properties

    return {
        "title": _clean(cp.title),
        "author": _clean(cp.author),
        "last_modified_by": _clean(cp.last_modified_by),
        "subject": _clean(cp.subject),
        "keywords": _clean(cp.keywords),
        "comments": _clean(cp.comments),
        "category": _clean(cp.category),
        "created": _clean(cp.created),
        "modified": _clean(cp.modified),
        "revision": _clean(cp.revision),
    }


def _extract_xlsx(content: bytes) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    props = wb.properties

    return {
        "title": _clean(props.title),
        "creator": _clean(props.creator),
        "last_modified_by": _clean(props.lastModifiedBy),
        "subject": _clean(props.subject),
        "description": _clean(props.description),
        "keywords": _clean(props.keywords),
        "category": _clean(props.category),
        "created": _clean(props.created),
        "modified": _clean(props.modified),
        "sheet_names": ", ".join(wb.sheetnames),
    }


def _convert_gps(value, ref) -> float:
    if not value or not ref:
        return None
    try:
        degrees, minutes, seconds = value
        decimal = float(degrees) + float(minutes) / 60 + float(seconds) / 3600
        if ref in ("S", "W"):
            decimal = -decimal
        return round(decimal, 6)
    except Exception:
        return None


def _extract_image(content: bytes) -> dict:
    from PIL import Image
    from PIL.ExifTags import TAGS, GPSTAGS

    img = Image.open(io.BytesIO(content))
    result = {
        "dimensions": f"{img.width}x{img.height}",
        "format": img.format,
        "color_mode": img.mode,
    }

    exif_raw = img.getexif()
    if not exif_raw:
        return result

    exif = {TAGS.get(k, k): v for k, v in exif_raw.items()}

    for key, label in (("Make", "camera_make"), ("Model", "camera_model"),
                        ("Software", "software"), ("DateTime", "date_taken"),
                        ("DateTimeOriginal", "date_taken_original")):
        if exif.get(key):
            result[label] = _clean(exif[key])

    # GPS lives in a nested IFD (tag 0x8825 / "GPSInfo")
    gps_ifd = exif_raw.get_ifd(0x8825) if hasattr(exif_raw, "get_ifd") else exif.get("GPSInfo")
    if gps_ifd:
        gps = {GPSTAGS.get(k, k): v for k, v in gps_ifd.items()}
        lat = _convert_gps(gps.get("GPSLatitude"), gps.get("GPSLatitudeRef"))
        lon = _convert_gps(gps.get("GPSLongitude"), gps.get("GPSLongitudeRef"))
        if lat is not None and lon is not None:
            result["gps_latitude"] = lat
            result["gps_longitude"] = lon
            result["gps_maps_link"] = f"https://www.google.com/maps?q={lat},{lon}"

    return result


def get_file_metadata(url: str) -> dict:
    """
    Master function — downloads a file from a public URL and extracts
    whatever metadata it can find based on detected file type.
    """
    if not url or not url.lower().startswith(("http://", "https://")):
        return {"error": "Please provide a valid http(s) URL"}

    headers = {"User-Agent": "DeepFinder/1.0 (file metadata OSINT scanner)"}

    try:
        resp = requests.get(url, headers=headers, timeout=DOWNLOAD_TIMEOUT, stream=True)
    except requests.exceptions.Timeout:
        return {"error": "Request timed out while downloading the file"}
    except requests.exceptions.ConnectionError:
        return {"error": "Could not connect to host"}
    except Exception as e:
        return {"error": str(e)}

    if resp.status_code != 200:
        return {"error": f"Server returned status {resp.status_code}"}

    content_type = resp.headers.get("Content-Type", "")

    content = b""
    try:
        for chunk in resp.iter_content(chunk_size=65536):
            content += chunk
            if len(content) > MAX_FILE_SIZE:
                return {"error": f"File exceeds the {MAX_FILE_SIZE // (1024*1024)}MB scan limit"}
    except Exception as e:
        return {"error": f"Download failed: {e}"}

    if not content:
        return {"error": "Downloaded file is empty"}

    file_type = _detect_type(url, content_type, content)

    http_meta = {
        "content_type": _clean(content_type),
        "size": _human_size(len(content)),
        "last_modified_header": _clean(resp.headers.get("Last-Modified")),
        "server": _clean(resp.headers.get("Server")),
    }

    try:
        if file_type == "pdf":
            extracted = _extract_pdf(content)
        elif file_type == "docx":
            extracted = _extract_docx(content)
        elif file_type == "xlsx":
            extracted = _extract_xlsx(content)
        elif file_type in IMAGE_EXTS:
            extracted = _extract_image(content)
        else:
            extracted = {}
    except Exception as e:
        extracted = {"parse_error": f"Could not parse {file_type} structure: {e}"}

    return {
        "url": url,
        "file_type": file_type,
        "metadata": {**http_meta, **extracted}
    }